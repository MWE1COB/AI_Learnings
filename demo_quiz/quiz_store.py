"""Excel-backed storage for demo quiz attempts (admin dashboard + export)."""
import io
import os
import sys
import threading
from datetime import datetime
import msoffcrypto
from openpyxl import Workbook, load_workbook
from msoffcrypto.format.ooxml import OOXMLFile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from runtime_paths import app_dir, secure_file  # noqa: E402
import admin_secrets  # noqa: E402

EXCEL_PATH = os.path.join(app_dir(), "demo_results.xlsx")

HEADERS = ["Name", "NTID", "Topic", "Level", "Score", "Total Questions",
           "Percentage", "Duration (sec)", "Submitted At"]

_lock = threading.Lock()


def _storage_password():
    """Password protecting the on-disk workbook — same as the export password, so
    copying demo_results.xlsx off the machine doesn't leak raw results either.
    """
    return admin_secrets.get_export_password()


def _read_workbook():
    """Load the on-disk workbook, transparently decrypting it if needed. Returns
    None if the file doesn't exist yet.
    """
    if not os.path.exists(EXCEL_PATH):
        return None
    with open(EXCEL_PATH, "rb") as f:
        raw = f.read()
    office_file = msoffcrypto.OfficeFile(io.BytesIO(raw))
    if not office_file.is_encrypted():
        return load_workbook(io.BytesIO(raw))  # legacy plaintext file predating at-rest encryption
    password = _storage_password()
    if not password:
        raise RuntimeError("demo_results.xlsx is encrypted but no export password is configured on this machine.")
    office_file.load_key(password=password)
    decrypted = io.BytesIO()
    office_file.decrypt(decrypted)
    decrypted.seek(0)
    return load_workbook(decrypted)


def _write_workbook(wb):
    """Save the workbook to disk, encrypted at rest with the export password
    when one is configured (so opening the file outside the app requires it).
    """
    plain = io.BytesIO()
    wb.save(plain)
    plain.seek(0)
    password = _storage_password()
    if password:
        encrypted = io.BytesIO()
        OOXMLFile(plain).encrypt(password, encrypted)
        data = encrypted.getvalue()
    else:
        data = plain.getvalue()
    with open(EXCEL_PATH, "wb") as f:
        f.write(data)
    secure_file(EXCEL_PATH)


def _ensure_workbook():
    with _lock:
        wb = _read_workbook()
        if wb is None:
            wb = Workbook()
            ws = wb.active
            ws.title = "Results"
            ws.append(HEADERS)
            _write_workbook(wb)
            return

        # Migrate older workbooks saved with a "Phone" column to the current schema.
        ws = wb["Results"]
        header = [c.value for c in ws[1]]
        if header == HEADERS:
            return
        if "Phone" in header:
            phone_idx = header.index("Phone")
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            ws.delete_rows(1, ws.max_row)
            ws.append(HEADERS)
            for row in rows:
                ws.append([v for i, v in enumerate(row) if i != phone_idx])
            _write_workbook(wb)


def save_attempt(name, ntid, topic, level_name, score, total, duration_sec):
    """Append a completed quiz attempt to the Excel workbook. Returns the saved record."""
    _ensure_workbook()
    percentage = round((score / total) * 100, 1) if total else 0
    submitted_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    row = [name, ntid, topic, level_name, score, total, percentage, duration_sec, submitted_at]

    with _lock:
        wb = _read_workbook()
        ws = wb["Results"]
        ws.append(row)
        _write_workbook(wb)

    return dict(zip(
        ["name", "ntid", "topic", "level", "score", "total", "percentage", "duration_sec", "submitted_at"],
        row,
    ))


def get_all_attempts():
    """Return all saved attempts (most recent first) for the admin dashboard."""
    _ensure_workbook()
    with _lock:
        wb = _read_workbook()
        ws = wb["Results"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))

    keys = ["name", "ntid", "topic", "level", "score", "total", "percentage", "duration_sec", "submitted_at"]
    attempts = [dict(zip(keys, row)) for row in rows if row and row[0]]
    attempts.reverse()
    return attempts


def ntid_has_attempted(ntid):
    """Return True if the NTID already has a recorded attempt."""
    _ensure_workbook()
    with _lock:
        wb = _read_workbook()
        ws = wb["Results"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and str(row[1]).strip().lower() == ntid.strip().lower():
                return True
    return False


def delete_attempt_by_ntid(ntid):
    """Remove all rows for the given NTID (allows a re-attempt)."""
    _ensure_workbook()
    with _lock:
        wb = _read_workbook()
        ws = wb["Results"]
        rows_to_keep = [list(ws[1])]  # keep header
        for row in ws.iter_rows(min_row=2):
            if str(row[1].value).strip().lower() != ntid.strip().lower():
                rows_to_keep.append([cell.value for cell in row])
        ws.delete_rows(1, ws.max_row)
        for row in rows_to_keep:
            ws.append([c if isinstance(c, (str, int, float, type(None))) else c for c in row])
        _write_workbook(wb)


def export_encrypted(password):
    """Return a workbook matching the dashboard view, password-protected for download."""
    _ensure_workbook()
    with _lock:
        raw_wb = _read_workbook()

    raw_ws = raw_wb["Results"]
    raw_header = [c.value for c in raw_ws[1]]

    # Build export workbook with dashboard-aligned columns.
    export_wb = Workbook()
    export_ws = export_wb.active
    export_ws.title = "Results"
    export_headers = ["#", "Name", "NTID", "Topic", "Level", "Score", "Percentage", "Duration (mm:ss)", "Submitted At"]
    export_ws.append(export_headers)

    rows = list(raw_ws.iter_rows(min_row=2, values_only=True))
    rows = [r for r in rows if r and r[0]]

    def _get(row, col):
        try:
            idx = raw_header.index(col)
            return row[idx]
        except (ValueError, IndexError):
            return ""

    for rank, row in enumerate(rows, start=1):
        score = _get(row, "Score")
        total = _get(row, "Total Questions")
        dur = _get(row, "Duration (sec)") or 0
        score_str = f"{score}/{total}" if total else str(score)
        dur_str = f"{int(dur) // 60}:{int(dur) % 60:02d}"
        export_ws.append([
            rank,
            _get(row, "Name"),
            _get(row, "NTID"),
            _get(row, "Topic"),
            _get(row, "Level"),
            score_str,
            _get(row, "Percentage"),
            dur_str,
            _get(row, "Submitted At"),
        ])

    plain = io.BytesIO()
    export_wb.save(plain)
    plain.seek(0)
    encrypted = io.BytesIO()
    OOXMLFile(plain).encrypt(password, encrypted)
    encrypted.seek(0)
    return encrypted
